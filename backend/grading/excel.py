# pyright: basic
"""파이프라인 P5 엑셀 출력: 채점 결과 목록 → .xlsx (로드맵 2-4-3, openpyxl).

[입력 → 출력]
  [grade, grade, ...](P4 출력, 답안지별) + AnswerKey + exam ──write_results_xlsx()──► 엑셀 파일

[시트 구성]
  1. 학생별 점수표 : 학번, 상태(확정 / 검수 대기 / 재촬영 필요), 확정 점수, 최고 가능 점수, 정답·오답 수
  2. 문항별 결과   : 학생 × 문항 표. 칸에 읽은 답, 색으로 정답/오답/무응답/복수/검수 대기 구분
  3. 문항별 정답률 : 문항별 정답률과 선지별 선택 수 (확정된 응답만 집계)
  4. 검수 대기     : 사람이 확인해야 할 문항·학번 목록과 칸별 채움 비율
  5. 정답표        : 채점에 쓴 정답과 배점

[확정 vs 검수 대기 구분] (지시서: 검수 대기 결과와 확정 점수는 구분)
  - '확정 점수' 열: 확정된 답안지에만 값
  - 검수 대기 답안지: '현재 확정분' ~ '최고 가능 점수'로 범위 표시

[색 규칙] 초록=정답, 빨강=오답, 주황=복수 응답, 회색=무응답, 노랑=검수 대기

[records(답안지별 grade dict)에서 쓰는 키]
  score.grade_sheet() / score.retake_record()가 만든 키를 그대로 씀
    status, student_id, student_id_raw, student_id_ok, score, confirmed_score, max_possible,
    total_points, pending_questions, counts, items(문항별: q, answer, result, source)
  호출하는 쪽(app/grading_service.py)이 덧붙이는 키
    source : 사진 파일 이름 (정렬·'파일' 열 표시용)
    read   : P3 판독 결과 원본 (있으면 '검수 대기' 시트에 칸별 채움 비율 표시)
    error  : 재촬영 필요 사유 (retake_record()가 넣음, '비고' 열에 표시)

[다른 파일과의 관계]
  - 이 파일을 부르는 쪽: pipeline/flow.py의 step_export (app/grading_service.py의 export_results가 이를 거쳐 호출)
  - 이 파일이 부르는 쪽: answer_key.py의 AnswerKey(정답·배점), score.py의 상태 상수, openpyxl(엑셀 쓰기)

[주요 상수·함수]
  - FILL / STATUS_FILL / RESULT_KO / BOX / CENTER : 칸 색·테두리·정렬 등 서식 상수
  - _ans_text()         : 답 목록 → 칸에 쓸 글자 ("3", "2,5", "-", "?")
  - _header()           : 머리글 행 한 줄 쓰기 (굵게·하늘색·테두리)
  - _widths()           : 열 너비 한꺼번에 지정
  - _sort_key()         : 답안지 정렬 기준 (상태 → 학번 → 파일명)
  - write_results_xlsx(): 엑셀 생성 메인 함수 (시트 5개 작성 후 저장)

[사용 예]
    records = [grade_sheet(read1, key), retake_record("IMG_0003.jpg", "정렬 실패")]
    write_results_xlsx("storage/exams/EX-.../results.xlsx", records, key,
                       exam={"exam_name": "중간고사", "subject": "컴퓨터비전", "exam_id": "EX-..."})
"""
from __future__ import annotations

import os
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .answer_key import AnswerKey
from .score import STATUS_CONFIRMED, STATUS_PENDING, STATUS_RETAKE

# 칸 배경색 (엑셀 연한 기본 색 계열). 키 = 문항 결과 이름 + 머리글/정답 행용
FILL = {
    "correct": PatternFill("solid", fgColor="C6EFCE"),
    "wrong": PatternFill("solid", fgColor="FFC7CE"),
    "multi": PatternFill("solid", fgColor="F4B084"),
    "blank": PatternFill("solid", fgColor="E7E6E6"),
    "pending": PatternFill("solid", fgColor="FFEB9C"),
    "head": PatternFill("solid", fgColor="DDEBF7"),   # 머리글 행 (하늘색)
    "key": PatternFill("solid", fgColor="FFF2CC"),    # '문항별 결과' 시트의 정답 행 (연노랑)
}
# 답안지 상태 칸 색: 확정=초록, 검수 대기=노랑, 재촬영 필요=빨강
STATUS_FILL = {STATUS_CONFIRMED: FILL["correct"], STATUS_PENDING: FILL["pending"], STATUS_RETAKE: FILL["wrong"]}
# 문항 결과 영어 이름 → 엑셀에 쓰는 한글 이름 (범례 표시용)
RESULT_KO = {"correct": "정답", "wrong": "오답", "multi": "복수 응답", "blank": "무응답", "pending": "검수 대기"}
THIN = Side(style="thin", color="BFBFBF")                       # 얇은 회색 선
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)      # 칸 네 변 테두리
CENTER = Alignment(horizontal="center", vertical="center")      # 가로·세로 가운데 정렬


def _ans_text(answer: list[int] | None) -> str:
    """답 목록 → 엑셀 칸에 쓸 짧은 글자.

    매개변수:
        answer: 답 목록 ([3] / [2, 5] / [] / None)
    반환:
        str. None → "?"(판정 보류), [] → "-"(무응답), 그 밖 → 쉼표로 이음 (예: [2, 5] → "2,5")
    """
    if answer is None:
        return "?"
    if not answer:
        return "-"
    return ",".join(str(a) for a in answer)


def _header(ws: Worksheet, row: int, values: list[Any]) -> None:
    """머리글 행 한 줄 쓰기. A열부터 values를 차례로 넣고 굵게·하늘색·가운데·테두리 서식 적용.

    매개변수:
        ws: 쓸 시트
        row: 행 번호 (1부터, 예: 4)
        values: 머리글 글자 목록 (예: ["번호", "학번", "상태"])
    반환:
        없음
    """
    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c, v)
        cell.font = Font(bold=True)
        cell.fill = FILL["head"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def _widths(ws: Worksheet, widths: dict[int, float]) -> None:
    """열 너비 한꺼번에 지정.

    매개변수:
        ws: 시트
        widths: {열 번호(1부터): 너비(엑셀 글자 폭 단위)} (예: {1: 6, 2: 14} → A열 6, B열 14)
    반환:
        없음
    """
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w


def _sort_key(r: dict[str, Any]) -> tuple[int, str, str]:
    """답안지 정렬 기준: 상태(확정 → 검수 대기 → 재촬영 필요) → 학번 → 파일명.

    학번이 없으면 읽은 원래 학번(student_id_raw), 그것도 없으면 "~"를 써서 뒤쪽으로 보냄.

    매개변수:
        r: 답안지 하나의 grade dict
    반환:
        (상태 순서 0~3, 학번 문자열, 파일명) 튜플 (예: (0, "20261234", "IMG_0001.jpg"))
    """
    order = {STATUS_CONFIRMED: 0, STATUS_PENDING: 1, STATUS_RETAKE: 2}
    return order.get(r["status"], 3), r.get("student_id") or r.get("student_id_raw") or "~", r.get("source", "")


def write_results_xlsx(path: str, records: list[dict[str, Any]], key: AnswerKey,
                       exam: dict[str, Any] | None = None) -> str:
    """엑셀 생성 메인 함수.
    records: grade_sheet() 결과 목록. 각 항목에 "source"(파일명), 선택적으로 "read"(판독 결과)를 붙여서 넘김
             ("read"가 있으면 '검수 대기' 시트에 칸별 채움 비율이 나옴)
    exam   : 시험 정보 (제목줄에 시험명·과목명 표시용, 없어도 됨)

    답안지를 _sort_key() 순서로 정렬한 뒤 시트 5개(학생별 점수표, 문항별 결과, 문항별 정답률,
    검수 대기, 정답표)를 차례로 만들고 path에 저장함. 저장 폴더가 없으면 만듦.
    '문항별 결과'와 '문항별 정답률'에는 재촬영 필요 답안지를 빼고(graded) 넣음.

    매개변수:
        path: 저장할 엑셀 경로 (예: "결과.xlsx")
        records: 답안지별 grade dict 목록 (retake_record() 결과를 섞어도 됨, 모듈 설명 참고)
        key: 채점에 쓴 정답표
        exam: 시험 정보 dict. 쓰는 키 exam_name, subject, exam_id (모두 선택, None 가능)
    반환:
        저장한 파일 경로 (path 그대로)
    예외:
        KeyError: grade dict에 "status" 같은 필수 키가 없을 때
    """
    exam = exam or {}
    records = sorted(records, key=_sort_key)                            # 상태 → 학번 → 파일명 순
    graded = [r for r in records if r["status"] != STATUS_RETAKE]       # 채점된 답안지만 (재촬영 필요 제외)
    wb = Workbook()

    # ---------- 1. 학생별 점수표 ----------
    ws = wb.active
    assert isinstance(ws, Worksheet)
    ws.title = "학생별 점수표"
    # 1행: 제목 (시험명 · 과목명 · 시험 ID 중 있는 것만 이어 붙임), 2행: 생성 시각과 답안지 수 요약
    title = " · ".join(v for v in (exam.get("exam_name"), exam.get("subject"), exam.get("exam_id") or key.exam_id) if v)
    ws["A1"] = f"채점 결과  {title}".strip()
    ws["A1"].font = Font(bold=True, size=13)
    n_conf = sum(r["status"] == STATUS_CONFIRMED for r in records)
    n_pend = sum(r["status"] == STATUS_PENDING for r in records)
    n_ret = sum(r["status"] == STATUS_RETAKE for r in records)
    ws["A2"] = (f"생성 {datetime.now():%Y-%m-%d %H:%M} · 답안지 {len(records)}장 "
                f"(확정 {n_conf}, 검수 대기 {n_pend}, 재촬영 필요 {n_ret}) · 만점 {key.total_points:g}점")
    # 4행: 머리글, 5행부터 답안지 한 장당 한 줄
    cols = ["번호", "학번", "상태", "확정 점수", "현재 확정분", "최고 가능 점수", "만점",
            "정답", "오답", "복수 응답", "무응답", "검수 대기 문항", "파일", "비고"]
    _header(ws, 4, cols)
    for i, r in enumerate(records, 1):
        c = r.get("counts", {})       # 결과 종류별 문항 수 (재촬영 필요면 빈 dict)
        # 비고: 재촬영 사유(error)가 있으면 그것, 채점된 답안지인데 학번이 불완전하면 '학번 확인 필요'
        note = r.get("error") or ("학번 확인 필요 (읽은 값: " + str(r.get("student_id_raw")) + ")"
                                  if r["status"] != STATUS_RETAKE and not r.get("student_id_ok") else "")
        row = [i, r.get("student_id") or "", r["status"], r.get("score"),
               r.get("confirmed_score") if r["status"] != STATUS_RETAKE else None,
               r.get("max_possible"), r.get("total_points"),
               c.get("correct"), c.get("wrong"), c.get("multi"), c.get("blank"),
               ", ".join(map(str, r.get("pending_questions", []))), r.get("source", ""), note]
        ws.append(row)
        rr = ws.max_row               # 방금 추가한 행 번호
        # 테두리는 전체 열, 가운데 정렬은 1~12열(번호 ~ 검수 대기 문항)만. 상태 칸은 상태별 색
        for col in range(1, len(cols) + 1):
            ws.cell(rr, col).border = BOX
            if col <= 12:
                ws.cell(rr, col).alignment = CENTER
        ws.cell(rr, 3).fill = STATUS_FILL.get(r["status"], FILL["blank"])
    # 맨 아래 요약 줄: 확정된 점수만으로 평균·최고·최저 (확정 답안지가 없으면 "-")
    if graded:
        conf_scores = [r["score"] for r in graded if r["score"] is not None]
        ws.append([])
        ws.append(["", "확정 평균", f"{sum(conf_scores) / len(conf_scores):.2f}" if conf_scores else "-",
                   "최고", max(conf_scores) if conf_scores else "-", "최저", min(conf_scores) if conf_scores else "-"])
        for col in range(2, 8):
            ws.cell(ws.max_row, col).font = Font(bold=True)
    _widths(ws, {1: 6, 2: 14, 3: 12, 4: 10, 5: 11, 6: 13, 7: 7, 8: 6, 9: 6, 10: 9, 11: 7, 12: 16, 13: 28, 14: 34})
    ws.freeze_panes = "C5"            # 머리글(1~4행)과 번호·학번 열(A~B) 고정

    # ---------- 2. 문항별 결과 ----------
    ws2 = wb.create_sheet("문항별 결과")
    nq = key.num_questions            # 문항 수
    # 1행: 머리글 (학번 | 상태 | 점수 | 1번 | 2번 | ...), 2행: 정답 행 (연노랑)
    _header(ws2, 1, ["학번", "상태", "점수"] + [f"{q}번" for q in range(1, nq + 1)])
    ws2.append(["정답", "", key.total_points] + [_ans_text(key.answers[q]) for q in range(1, nq + 1)])
    for col in range(1, nq + 4):
        ws2.cell(2, col).fill = FILL["key"]
        ws2.cell(2, col).font = Font(bold=True)
        ws2.cell(2, col).alignment = CENTER
        ws2.cell(2, col).border = BOX
    # 3행부터 답안지 한 장당 한 줄. 점수는 확정이면 숫자, 검수 대기면 "현재 확정분~최고 가능" 범위 글자
    for r in graded:
        score = r["score"] if r["score"] is not None else f"{r['confirmed_score']:g}~{r['max_possible']:g}"
        ws2.append([r.get("student_id") or r.get("student_id_raw"), r["status"], score]
                   + [_ans_text(it["answer"]) for it in r["items"]])
        rr = ws2.max_row
        for col in range(1, 4):
            ws2.cell(rr, col).alignment = CENTER
            ws2.cell(rr, col).border = BOX
        ws2.cell(rr, 2).fill = STATUS_FILL.get(r["status"], FILL["blank"])
        # 문항 칸: 결과별 색, 검수에서 고친 문항은 굵은 밑줄
        for j, it in enumerate(r["items"]):
            cell = ws2.cell(rr, 4 + j)
            cell.fill = FILL[it["result"]]
            cell.alignment = CENTER
            cell.border = BOX
            if it["source"] == "review":
                cell.font = Font(bold=True, underline="single")
    # 표 아래 한 줄 띄고 범례 (결과별 색 견본 + 칸 표기 설명)
    lr = ws2.max_row + 2              # 범례 시작 행
    ws2.cell(lr, 1, "범례").font = Font(bold=True)
    for j, (k, label) in enumerate(RESULT_KO.items(), 1):
        cell = ws2.cell(lr + j, 1, label)
        cell.fill = FILL[k]
        cell.alignment = CENTER
    ws2.cell(lr + len(RESULT_KO) + 1, 1,
             "칸 표기: 숫자=읽은 답, '-'=무응답, '?'=판정 보류, '2,5'=복수 응답, 굵은 밑줄=검수에서 수정")
    _widths(ws2, {1: 14, 2: 11, 3: 9, **{c: 5.5 for c in range(4, nq + 4)}})
    ws2.freeze_panes = "D3"           # 머리글·정답 행(1~2행)과 학번·상태·점수 열(A~C) 고정

    # ---------- 3. 문항별 정답률 ----------
    ws3 = wb.create_sheet("문항별 정답률")
    k = key.num_choices               # 선지 수
    # 머리글: 문항 | 정답 | 배점 | 확정 응답 수 | 정답 수 | 정답률 | 1번 선택 ... k번 선택 | 복수 응답 | 무응답 | 검수 대기
    _header(ws3, 1, ["문항", "정답", "배점", "확정 응답 수", "정답 수", "정답률"]
            + [f"{c}번 선택" for c in range(1, k + 1)] + ["복수 응답", "무응답", "검수 대기"])
    for qn in range(1, nq + 1):
        its = [r["items"][qn - 1] for r in graded]                    # 이 문항의 답안지별 채점 기록
        decided = [it for it in its if it["result"] != "pending"]       # 확정된 응답만 (검수 대기 제외)
        n_ok = sum(it["result"] == "correct" for it in decided)         # 정답 수
        # 선지별 선택 수: 1개만 표시한 응답(correct/wrong) 중 c번을 고른 수 (복수·무응답은 제외)
        dist = [sum(it["result"] in ("correct", "wrong") and it["answer"] == [c] for it in decided)
                for c in range(1, k + 1)]
        ws3.append([qn, _ans_text(key.answers[qn]), key.point(qn), len(decided), n_ok,
                    n_ok / len(decided) if decided else None] + dist
                   + [sum(it["result"] == "multi" for it in its), sum(it["result"] == "blank" for it in its),
                      sum(it["result"] == "pending" for it in its)])
        rr = ws3.max_row
        ws3.cell(rr, 6).number_format = "0.0%"        # 정답률 칸을 백분율로 표시 (예: 0.85 → 85.0%)
        for col in range(1, 10 + k):
            ws3.cell(rr, col).alignment = CENTER
            ws3.cell(rr, col).border = BOX
        # 정답 선지의 'c번 선택' 칸(7열부터)을 초록으로 칠함
        for c in key.answers[qn]:
            if 1 <= c <= k:
                ws3.cell(rr, 6 + c).fill = FILL["correct"]
    ws3.cell(ws3.max_row + 2, 1, "정답률 = 정답 수 / 확정 응답 수 (검수 대기는 제외). 초록 칸 = 정답 선지")
    _widths(ws3, {1: 6, 2: 7, 3: 6, 4: 11, 5: 8, 6: 8, **{c: 8 for c in range(7, 10 + k)}})
    ws3.freeze_panes = "B2"           # 머리글 행과 문항 열 고정

    # ---------- 4. 검수 대기 ----------
    ws4 = wb.create_sheet("검수 대기")
    # 사람이 확인해야 할 항목을 한 줄씩: 재촬영 필요 답안지 / 불완전한 학번 / 판정 보류·복수 응답 문항
    _header(ws4, 1, ["파일", "학번(읽은 값)", "대상", "사유", "읽은 답", "칸별 채움 비율"])
    for r in records:
        if r["status"] == STATUS_RETAKE:
            ws4.append([r.get("source"), "", "답안지 전체", "재촬영 필요: " + str(r.get("error", "")), "", ""])
            continue
        read = r.get("read") or {}    # P3 판독 원본 (없으면 채움 비율 칸이 비게 됨)
        # 학번이 불완전하면: 정상(ok)이 아닌 자리만 "3자리 blank" 처럼 나열
        if not r.get("student_id_ok"):
            fills = "; ".join(f"{d['digit_index'] + 1}자리 {d['status']}" for d in read.get("student_id", {}).get("digits", [])
                              if d["status"] != "ok")
            ws4.append([r.get("source"), r.get("student_id_raw"), "학번", "학번 판독 불완전", r.get("student_id_raw"), fills])
        mcq = {q["q"]: q for q in read.get("mcq", [])}   # 문항 번호 → 판독 원본 한 문항
        # 판정 보류·복수 응답 문항: 선지별 채움 비율을 "1:0.12  2:0.85 ..." 형태로 함께 표시
        for it in r["items"]:
            if it["result"] in ("pending", "multi"):
                q = mcq.get(it["q"], {})
                fills = "  ".join(f"{c}:{v:.2f}" for c, v in enumerate(q.get("fill", []), 1))
                reason = "애매한 칸 (판정 보류)" if it["result"] == "pending" else "복수 응답 (0점 처리, 확인 권장)"
                ws4.append([r.get("source"), r.get("student_id") or r.get("student_id_raw"),
                            f"{it['q']}번", reason, _ans_text(it["answer"]), fills])
    for row in ws4.iter_rows(min_row=2):
        for cell in row:
            cell.border = BOX
    _widths(ws4, {1: 28, 2: 14, 3: 12, 4: 30, 5: 9, 6: 44})
    ws4.freeze_panes = "A2"           # 머리글 행 고정

    # ---------- 5. 정답표 ----------
    ws5 = wb.create_sheet("정답표")
    # 채점에 쓴 정답과 배점 기록 + 맨 아래 합계(만점)
    _header(ws5, 1, ["문항", "정답", "배점"])
    for qn in range(1, nq + 1):
        ws5.append([qn, _ans_text(key.answers[qn]), key.point(qn)])
        for col in range(1, 4):
            ws5.cell(ws5.max_row, col).alignment = CENTER
            ws5.cell(ws5.max_row, col).border = BOX
    ws5.append(["합계", "", key.total_points])
    _widths(ws5, {1: 8, 2: 8, 3: 8})

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
